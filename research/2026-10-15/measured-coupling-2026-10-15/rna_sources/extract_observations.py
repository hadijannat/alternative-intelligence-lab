"""Extract numerical measurements with workbook coordinates, never graph digitization."""
import csv
import hashlib
import json
from pathlib import Path
import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent

def main():
    workbook = ROOT / "raw/salibi_data.xlsx"
    wb = openpyxl.load_workbook(workbook, data_only=True)
    rows = []
    for sheet, row_range, x_label in [("Figure 2", range(5, 9), "ft_cycles"), ("Figure 3", range(4, 12), "generation")]:
        ws = wb[sheet]
        for r in row_range:
            for seed, start in [(0, 3), (1, 6), (2, 9)]:
                for rep in range(3):
                    c = start + rep
                    value = ws.cell(r, c).value
                    assert isinstance(value, (int, float)) and 0 <= value <= 100
                    rows.append(dict(dataset=sheet.replace(" ", "_"), condition_id="salibi_r3c_guv",
                                     seed_uM=seed, hyper_a_uM=10, x_name=x_label, x=ws.cell(r,2).value,
                                     replicate=rep+1, value_percent=value, sd_percent="", n=1,
                                     source_sheet=sheet, source_cell=ws.cell(r,c).coordinate))
    ws = wb["S13"]
    for kind, seeds, starts in [("F1", [.1, .25, .5, 1.], [6,9,12,15]), ("HyperA", [1.,1.5,2.,2.5], [20,23,26,29])]:
        for seed, c in zip(seeds, starts):
            for r in range(6,11):
                value = ws.cell(r,c).value
                assert isinstance(value, (int,float)) and 0 <= value <= 100
                rows.append(dict(dataset="S13_"+kind,condition_id="salibi_r3c_bulk_sucrose",
                                 seed_uM=seed if kind=="F1" else 0,
                                 hyper_a_uM=seed if kind=="HyperA" else 0,
                                 x_name="time_min",x=ws.cell(r,2 if kind=="F1" else 19).value,
                                 replicate="summary",value_percent=value,sd_percent=ws.cell(r,c+1).value,n=ws.cell(r,c+2).value,
                                 source_sheet="S13",source_cell=ws.cell(r,c).coordinate))
    with (ROOT/"observations.csv").open("w",newline="") as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0])); writer.writeheader();writer.writerows(rows)
    for name in ["ameta", "salibi"]:
        pages=PdfReader(ROOT/"raw"/(name+"_si.pdf")).pages
        (ROOT/"raw"/(name+"_si.txt")).write_text("\n\n".join(f"PDF PAGE {i+1}\n"+p.extract_text() for i,p in enumerate(pages)))
    summary={"rows":len(rows),"workbook_sha256":hashlib.sha256(workbook.read_bytes()).hexdigest(),
             "source":"Salibi et al. 2023, DOI 10.1038/s41467-023-36940-z, Source Data workbook",
             "license":"CC BY 4.0 article license; no separate notice identified for supplement",
             "transform":"Exact cell extraction; percent units retained. S13 summaries are not invented individual replicates. S13d duplicated endpoints excluded."}
    (ROOT/"extraction.json").write_text(json.dumps(summary,indent=2)+"\n")
    print(json.dumps(summary,indent=2))

if __name__ == "__main__": main()
