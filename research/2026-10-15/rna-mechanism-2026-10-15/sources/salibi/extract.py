"""Exact retained-workbook extraction and checks; no fitting or source mutation."""
import csv
import hashlib
import json
from pathlib import Path
import sys
import openpyxl

ROOT = Path(__file__).resolve().parent
OLD = ROOT.parents[2] / "measured-coupling-2026-10-15"
WORKBOOK = OLD / "rna_sources/raw/salibi_data.xlsx"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def extract():
    before = {str(p.relative_to(OLD)): sha(p) for p in OLD.rglob("*") if p.is_file()}
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    rows = []

    def add(sheet, cell, dataset, condition, measure, unit, x, x_unit, group,
            replicate="", sd_cell="", n_cell="", previously_used=False):
        ws = wb[sheet]
        value = ws[cell].value
        assert isinstance(value, (int, float)), (sheet, cell, value)
        rows.append(dict(dataset=dataset, condition_id=condition, measurement=measure,
                         value=value, unit=unit, x=x, x_unit=x_unit, group=group,
                         replicate=replicate, sd=ws[sd_cell].value if sd_cell else "",
                         n=ws[n_cell].value if n_cell else 1,
                         source_sheet=sheet, source_cell=cell,
                         source_sd_cell=sd_cell, source_n_cell=n_cell,
                         previously_used_for_model=previously_used))

    for sheet, rr, xunit in [("Figure 2", range(5, 9), "FT cycle"),
                              ("Figure 3", range(4, 12), "generation")]:
        for r in rr:
            for seed, c in [(0, 3), (1, 6), (2, 9)]:
                for rep in range(3):
                    add(sheet, wb[sheet].cell(r, c+rep).coordinate, sheet.replace(" ", "_"),
                        "guv_f1", "Cy5-F1 fraction", "percent", wb[sheet].cell(r, 2).value,
                        xunit, f"seed_{seed}_uM", rep+1, previously_used=True)
    for kind, seeds, cs in [("F1", [.1, .25, .5, 1.], [6, 9, 12, 15]),
                            ("HyperA", [1, 1.5, 2, 2.5], [20, 23, 26, 29])]:
        for seed, c in zip(seeds, cs):
            for r in range(6, 11):
                ws = wb["S13"]
                add("S13", ws.cell(r, c).coordinate, "S13_"+kind, "bulk_s13",
                    "Cy5-F1 fraction", "percent", ws.cell(r, 2 if kind=="F1" else 19).value,
                    "minute", f"{kind}_{seed}_uM", "summary",
                    ws.cell(r, c+1).coordinate, ws.cell(r, c+2).coordinate, True)
    for r in [6, 10]:
        add("S13", f"C{r}", "S13_negative", "bulk_s13", "Cy5-F1 fraction", "percent",
            wb["S13"][f"B{r}"].value, "minute", "no_seed", "summary", f"D{r}", f"E{r}")
    for r in range(5, 19):
        for group, c in [("reporter_only_seed1", 3), ("hyper9_reporter1_seed0", 6),
                          ("hyper9_reporter1_seed1", 9)]:
            for rep in range(3):
                add("S12", wb["S12"].cell(r, c+rep).coordinate, "S12", "bulk_s12",
                    "Cy5-F1 fraction", "percent", wb["S12"].cell(r, 2).value,
                    "published_generation_coordinate", group, rep+1)
    for r, fraction in [(6, 0), (7, .25), (8, .5)]:
        for group, c in [("no_seed", 3), ("nominal_seed1", 6)]:
            for rep in range(3):
                add("S15", wb["S15"].cell(r, c+rep).coordinate, "S15", "bulk_s15",
                    "Cy5-F1 fraction", "percent", fraction, "GUV suspension volume fraction",
                    group, rep+1)
    for r in range(5, 9):
        for group, c in [("seed0_FT_only", 3), ("seed0_FT_plus42C", 6),
                          ("seed1_FT_only", 9), ("seed1_FT_plus42C", 12)]:
            for rep in range(3):
                add("S18", wb["S18"].cell(r, c+rep).coordinate, "S18", "bulk_s18",
                    "Cy5-F1 fraction", "percent", wb["S18"].cell(r, 2).value,
                    "FT cycle", group, rep+1)
    for sheet, rr, measure, xunit, condition in [
        ("S9", range(6, 57), "GUV relative frequency", "fluorescence arbitrary units", "tracer_s9"),
        ("S17", range(6, 47), "GUV relative frequency", "diameter micrometre", "dye_s17")]:
        for r in rr:
            for c, group in [(3, "before_FT"), (4, "after_FT")]:
                add(sheet, wb[sheet].cell(r, c).coordinate, sheet+"_histogram", condition,
                    measure, "fraction", wb[sheet].cell(r, 2).value, xunit, group,
                    "published histogram")
    for cell, measure, group, unit in [("G5", "mean GUV fluorescence", "before_FT", "arbitrary units"),
                                        ("H5", "mean GUV fluorescence", "after_FT", "arbitrary units"),
                                        ("G6", "tracer content loss", "after_FT", "percent")]:
        add("S9", cell, "S9_summary", "tracer_s9", measure, unit, 1, "FT cycle", group,
            "published summary")
    with (ROOT/"observations.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)

    old_rows = list(csv.DictReader((OLD/"rna_sources/observations.csv").open()))
    for row in old_rows:
        assert float(row["value_percent"]) == wb[row["source_sheet"]][row["source_cell"]].value
    for row in rows:
        assert row["value"] == wb[row["source_sheet"]][row["source_cell"]].value
        if row["sd"] != "":
            assert row["sd"] == wb[row["source_sheet"]][row["source_sd_cell"]].value

    sys.path.insert(0, str(OLD/"rna_model"))
    from model import Parameters, bulk_curve
    p = Parameters(**json.loads((OLD/"rna_model/results-v1/summary.json").read_text())["parameters"])
    residual_rows = []
    for kind, seeds in [("F1", [.1, .25, .5, 1.]), ("HyperA", [1, 1.5, 2, 2.5])]:
        for seed in seeds:
            observed = [r for r in rows if r["dataset"] == "S13_"+kind and r["group"] == f"{kind}_{seed}_uM"]
            predicted = bulk_curve(p, seed if kind=="F1" else 0, seed if kind=="HyperA" else 0)
            for row, pred in zip(observed, predicted):
                residual_rows.append(dict(group=row["group"], time_min=row["x"],
                                          observed_percent=row["value"], predicted_percent=float(pred),
                                          error_pp=float(pred-row["value"]), source_cell=row["source_cell"]))
    with (ROOT/"v1_calibration_residuals.csv").open("w", newline="") as f:
        writer=csv.DictWriter(f, fieldnames=list(residual_rows[0]));writer.writeheader();writer.writerows(residual_rows)
    assert len(residual_rows)==40
    groups={g:(sum(r["error_pp"]**2 for r in residual_rows if r["group"]==g)/5)**.5
            for g in sorted({r["group"] for r in residual_rows})}
    assert all(sha(OLD/p)==value for p, value in before.items())
    result=dict(workbook_sha256=sha(WORKBOOK), exact_new_cell_checks=len(rows),
                prior_extraction_cell_checks=len(old_rows), old_files_unchanged=len(before),
                source_url="https://static-content.springer.com/esm/art%3A10.1038%2Fs41467-023-36940-z/MediaObjects/41467_2023_36940_MOESM8_ESM.xlsx",
                extraction="Exact numerical cells; no digitization; no model fitting; histogram bins are not paired lineages.",
                v1_calibration_rmse_by_condition_pp=groups,
                v1_total_calibration_rmse_pp=(sum(r["error_pp"]**2 for r in residual_rows)/40)**.5,
                tracer_s9_retained_intensity_fraction=wb["S9"]["H5"].value/wb["S9"]["G5"].value,
                all_checks_pass=True)
    (ROOT/"verification.json").write_text(json.dumps(result, indent=2)+"\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    extract()
