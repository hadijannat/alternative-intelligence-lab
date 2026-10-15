import hashlib
import io
import json
import math
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import openpyxl


BASE = Path(__file__).resolve().parent
NESTED = "41467_2021_26281_MOESM3_ESM.zip"
PINS = {
    "article.xml": "4c1460c430775bd3f44d80689aec4baa15fb0496464f1bacc59fe6706a79b265",
    "supplementary.zip": "15a142cb70b4fe9a8a1e8d00fb974cfedad178d3f46b67dae1879cf413181c78",
    NESTED: "231ddfd7f50be46e7ab4b62b9197270f7c4a72fc1293e8aca8f68397ec858d84",
}
POTENTIAL = "Potential (V vs. RHE)"
RATE = "Deposition rate (ng cm-2 min-1)"
RESCUE_HEADERS = {"A1": "Ni(II)", "C1": "Co(II)", "E1": "Fe(II)",
                  "A2": "Time (h)", "C2": "Time (h)", "E2": "Time (h)",
                  "B2": POTENTIAL, "D2": POTENTIAL, "F2": POTENTIAL}


def checked_bytes(data, name):
    if hashlib.sha256(data).hexdigest() != PINS[name]:
        raise ValueError(f"Source digest mismatch: {name}")
    return data


def number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Expected numeric source value, got {value!r}")
    if not math.isfinite(value):
        raise ValueError("Nonfinite source value")
    return value


def deposition_rows(sheet):
    if sheet.title != "Figure 1c" or sheet.max_column != 2:
        raise ValueError("Expected the two-column Figure 1c deposition sheet")
    rows = sheet.iter_rows()
    if [cell.value for cell in next(rows)] != [POTENTIAL, RATE]:
        raise ValueError("Unexpected deposition labels or units")
    result = []
    previous = -math.inf
    for potential, rate in rows:
        x, y = number(potential.value), number(rate.value)
        if x <= previous or y < 0:
            raise ValueError("Potentials must increase and rates must be nonnegative")
        previous = x
        result.append({
            "potential_V_vs_RHE": x,
            "deposition_rate_ng_cm-2_min-1": y,
            "source_cells": [potential.coordinate, rate.coordinate],
            "mean": None,
            "uncertainty": None,
            "replicate_count": None,
        })
    if not result:
        raise ValueError("No deposition observations")
    return result


def text_at(root, element_id):
    element = root.find(f".//*[@id='{element_id}']")
    if element is None:
        raise ValueError(f"Missing article element {element_id}")
    return " ".join("".join(element.itertext()).split())


def ion_addition_response(sheet):
    if sheet.title != "Figure 2e" or sheet.max_column != 6:
        raise ValueError("Expected six-column Figure 2e chronopotentiometry")
    rows = sheet.iter_rows()
    headers = {c.coordinate: c.value for _ in range(2) for c in next(rows) if c.value is not None}
    if headers != RESCUE_HEADERS:
        raise ValueError("Unexpected Figure 2e labels or units")
    series = {ion: {} for ion in ("Ni(II)", "Co(II)", "Fe(II)")}
    for row in rows:
        for offset, values in enumerate(series.values()):
            time, potential = row[2 * offset:2 * offset + 2]
            t, voltage = number(time.value), number(potential.value)
            if t < 0 or (values and t <= next(reversed(values))):
                raise ValueError("Times must be nonnegative and strictly increasing")
            values[t] = {"time_h": t, "potential_V_vs_RHE": voltage,
                         "source_cells": [time.coordinate, potential.coordinate]}
    common = set.intersection(*(set(values) for values in series.values()))
    if 0.5 not in common or max(common) <= 0.5:
        raise ValueError("Missing exact shared 0.5 h reference or later shared endpoint")
    endpoint = max(common)
    return {
        "selection_rule": "Exact recorded 0.5 h intervention-time point and latest exactly shared recorded time; no interpolation. The intervention-time sample is not known to precede addition.",
        "common_reference_time_h": 0.5, "common_endpoint_time_h": endpoint,
        "responses": [{"ion": ion, "reference": values[0.5], "endpoint": values[endpoint],
                       "endpoint_minus_reference_mV": round(1000 * (values[endpoint]["potential_V_vs_RHE"] - values[0.5]["potential_V_vs_RHE"]), 2),
                       "uncertainty": None, "replicate_count": None}
                      for ion, values in series.items()],
        "interpretation": "Descriptive within-curve potential changes. Different baselines and no supplied replicate uncertainty prevent a statistical treatment-effect claim; this is not deposition rate, atom replacement or target renewal.",
    }


def analyze():
    article = checked_bytes((BASE / "raw/article.xml").read_bytes(), "article.xml")
    outer = checked_bytes((BASE / "raw/supplementary.zip").read_bytes(), "supplementary.zip")
    root = ET.fromstring(article)
    fig1, fig2 = text_at(root, "Fig1"), text_at(root, "Fig2")
    if "c Deposition rates of Fe hydroxides on EQCM sensors" not in fig1:
        raise ValueError("Article no longer identifies Figure 1c as EQCM deposition")
    if "e Chronopotentiometric curves" not in fig2:
        raise ValueError("Article no longer identifies Figure 2e as chronopotentiometry")
    with zipfile.ZipFile(io.BytesIO(outer)) as archive:
        nested = checked_bytes(archive.read(NESTED), NESTED)
        outer_members = archive.namelist()
    workbooks = []
    observations = None
    rescue_headers = None
    rescue_response = None
    conflicting_label = None
    with zipfile.ZipFile(io.BytesIO(nested)) as archive:
        members = archive.namelist()
        for name in members:
            if not name.endswith(".xlsx"):
                continue
            data = archive.read(name)
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = []
            for sheet in book:
                headers = {
                    cell.coordinate: cell.value
                    for row in sheet.iter_rows(max_row=min(3, sheet.max_row))
                    for cell in row if isinstance(cell.value, str)
                }
                sheets.append({
                    "name": sheet.title,
                    "declared_rows_including_headers": sheet.max_row,
                    "declared_columns": sheet.max_column,
                    "text_cells_in_first_three_rows": headers,
                })
                if name.endswith("Figure_1.xlsx") and sheet.title == "Figure 1c":
                    observations = deposition_rows(sheet)
                if name.endswith("Figure_2.xlsx") and sheet.title == "Figure 2e":
                    rescue_headers = headers
                    rescue_response = ion_addition_response(sheet)
                if name.endswith("Figure_2.xlsx") and sheet.title == "Figure 2d":
                    conflicting_label = headers.get("C1")
            workbooks.append({"member": name, "bytes": len(data),
                              "sha256": hashlib.sha256(data).hexdigest(), "sheets": sheets})
            book.close()
    if observations is None or rescue_headers is None or conflicting_label != "with 50 mM Fe(II)":
        raise ValueError("Expected source sheets or original conflicting label missing")
    source = {"doi": "10.1038/s41467-021-26281-0", "pins_sha256": PINS,
              "nested_archive_member": NESTED, "openpyxl_version": openpyxl.__version__}
    inventory = {"source": source, "outer_members": outer_members,
                 "source_archive_members": members, "workbooks": workbooks,
                 "inspection_scope": "All workbook sheet names, dimensions and text cells in first three rows; all Figure 1c rows extracted; all Figure 2e rows validated and exact intervention-time/shared-endpoint samples extracted."}
    result = {
        "status": "SOURCE_DATA_RECOVERED_MATCHED_EQCM_COMPARATOR_UNAVAILABLE",
        "source": source,
        "panel_mapping": {"Fig1": {"article_locator": "//*[@id='Fig1']", "caption": fig1},
                          "Fig2": {"article_locator": "//*[@id='Fig2']", "caption": fig2}},
        "fig1c": {
            "workbook": "Soruce_Data/Soruce_Data_Figure_1.xlsx", "sheet": "Figure 1c",
            "units_source_cells": {"A1": POTENTIAL, "B1": RATE},
            "observations": observations,
            "interpretation": "Published numerical deposition-rate points. Source sheet does not label means, replicates or uncertainty; null is not zero. Exact zero entries are reported values, not a detection-limit claim.",
            "conditions": {"article_locator": "//*[@id='Par6']", "medium": "Fe-containing KBi electrolyte",
                           "aging_h_approximate": 100, "nominal_Fe_II_added_uM": 50,
                           "soluble_Fe_fresh_uM": 22.7, "soluble_Fe_aged_uM": 7.5,
                           "uncertainty": None},
        },
        "fig2e": {
            "workbook": "Soruce_Data/Soruce_Data_Figure_2.xlsx", "sheet": "Figure 2e",
            "source_header_cells": rescue_headers,
            "article_locator": "//*[@id='Fig2']/caption",
            "measured_quantity": "Potential versus time; ion-addition chronopotentiometry",
            "caption_conditions": {"pH": 14, "ion_addition_time_h": 0.5, "ion_addition_uM": 30},
            "current_density_mA_cm-2": None,
            "current_density_note": "Chronopotentiometric protocol; this panel's caption does not independently specify current density.",
            "deposition_rate": None,
            "descriptive_potential_response": rescue_response,
        },
        "matched_Co_assisted_vs_unassisted_EQCM": {
            "common_potential_comparisons": None,
            "Co_assisted_rate_data": None,
            "reason": "Figure 2e has time and potential axes, not deposition rate. The five-workbook inventory contains no second deposition-rate sheet. Article Par9 instead discusses sequential electrodeposition Figure 2b and Supplementary Figure S5. No interpolation, fitting or substituted current-to-deposition conversion was performed.",
        },
        "source_label_conflict": {
            "workbook": "Soruce_Data/Soruce_Data_Figure_2.xlsx", "sheet": "Figure 2d", "cell": "C1",
            "workbook_label_verbatim": conflicting_label,
            "article_locator": "//*[@id='Par10']", "article_nominal_feed_uM": 50,
            "disposition": "Unresolved mM versus microM discrepancy retained; no quantitative feed-response comparison uses this label.",
        },
        "separate_text_observations": {
            "no_added_Fe_film_loss": {
                "article_locator": "//*[@id='Par10']", "method": "ICP-MS",
                "initial_Fe_nmol_cm-2": 10.4, "final_Fe_nmol_cm-2": 5.4,
                "elapsed_h": 10, "current_mA_cm-2": 10, "medium": "0.25 M KBi, pH 14, no added Fe(II)",
                "net_loss_nmol_cm-2": 10.4 - 5.4, "uncertainty": None,
                "inference_limit": "Net film inventory loss in this no-added-feed experiment; neither gross dissolution nor founder replacement in a fed experiment.",
            },
            "KBi_soluble_inventory": {
                "article_locator": "//*[@id='Par6']", "nominal_Fe_II_added_uM": 50,
                "fresh_soluble_Fe_uM": 22.7, "aged_soluble_Fe_uM": 7.5, "aging_h_approximate": 100,
                "uncertainty": None,
                "inference_limit": "Nominal addition differs from soluble Fe inventory. These totals do not measure atom replacement, turnover flux or founder survival.",
            },
        },
        "target_renewal_demonstrated": False,
        "limitations": ["Retrospective descriptive source audit; no physical experiment or held-out test.",
                        "No isotope-resolved founder accounting or quantitative replacement fraction.",
                        "No transfer of alkaline driven-electrode behavior to the acidic Mo36-derived target hypothesis; actual target operating medium remains unverified.",
                        "Source-supported electrode function and proposed mechanism do not establish integrated target renewal."],
    }
    return result, inventory


if __name__ == "__main__":
    result, inventory = analyze()
    for name, value in [("analysis.json", result), ("source-data-inventory.json", inventory)]:
        (BASE / name).write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n")
    print(f"{result['status']}: {len(result['fig1c']['observations'])} deposition rows; {len(inventory['workbooks'])} workbooks")
