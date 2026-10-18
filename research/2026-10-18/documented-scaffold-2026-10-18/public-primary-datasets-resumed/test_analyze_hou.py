import json
from pathlib import Path
import tempfile
import unittest
from unittest import mock

import openpyxl

import analyze_hou as hou


class HouAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "synthetic.xlsx"

    def make_workbook(self, change=None):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Fig. 5e"
        sheet["A1"] = "Fig. 5e"
        for index, row in enumerate((3, 6, 9), start=1):
            sheet.cell(row, 1, f"origina-{index}")
            sheet.cell(row + 1, 1, f"selfhealing-{index}")
            sheet.cell(row, 2, 2 * index)
            sheet.cell(row + 1, 2, 3 * index)
        sheet["B5"] = sheet["B8"] = "--"
        for name, count in hou.TRACE_SHEETS.items():
            sheet = workbook.create_sheet(name)
            sheet["A1"] = f"Fig. S{name.rsplit(' ', 1)[1]}"
            for index in range(count):
                column = 1 + 3 * index
                sheet.cell(4, column, "time " if name.endswith("16") and column == 7 else "time")
                sheet.cell(4, column + 1, "current density")
                for row, (t, y) in enumerate(((20, -10), (50, -7), (120, 0)), start=5):
                    sheet.cell(row, column, t)
                    sheet.cell(row, column + 1, y)
            sheet.cell(9, 3, "tail dimension marker outside admitted pairs")
        if change:
            change(workbook)
        workbook.save(self.path)
        workbook.close()
        return hou.sha256(self.path)

    def test_linear_trace_irregular_grid_and_time_unit_invariance(self):
        pairs = [(12, -10), (15, -9.7), (49, -6.3), (112, 0)]
        record = hou.trace_record("Supplementary Fig. 16", 1, pairs)
        self.assertEqual(record["status"], "ok")
        self.assertAlmostEqual(record["early_mean_magnitude_native"], 8.5)
        self.assertAlmostEqual(record["late_mean_magnitude_native"], 0.5)
        self.assertAlmostEqual(record["retention"], 1 / 17)
        scaled = hou.trace_record("Supplementary Fig. 16", 4, [(t * 60, y) for t, y in pairs])
        self.assertAlmostEqual(record["retention"], scaled["retention"])
        self.assertAlmostEqual(record["reference_residual"], -16 / 17)

    def test_piecewise_kink_requires_area_integration(self):
        u = [0, 0.12, 0.18, 0.4, 1]
        magnitude = [2, 2, 8, 8, 8]
        self.assertAlmostEqual(hou.window_mean(u, magnitude, 0.1, 0.2), 5)
        self.assertAlmostEqual(hou.window_mean(u, magnitude, 0.9, 1), 8)

    def test_sign_changes_use_sampled_magnitudes_before_interpolation(self):
        record = hou.trace_record("Supplementary Fig. 18", 1, [(0, -4), (0.15, 4), (1, -4)])
        self.assertEqual(record["status"], "ok")
        self.assertAlmostEqual(record["early_mean_magnitude_native"], 4)
        self.assertAlmostEqual(record["late_mean_magnitude_native"], 4)
        self.assertAlmostEqual(record["retention"], 1)
        self.assertFalse(record["nominal_deterioration"])

    def test_malformed_traces_are_refused_without_partial_estimates(self):
        for pairs in (
            [(0, 1), (None, None), (1, 1)],
            [(0, 1), (1, None)], [(0, None), (1, 1)],
            [(0, 1), (1, "2")], [(0, 1), (1, "=2")],
            [(0, 1), (1, float("nan"))], [(0, 1), (float("inf"), 1)],
            [(0, 1), (0, 2)], [(1, 1), (0, 2)],
            [(0, 1)], [], [(0, 0), (1, 0)], [(0, 1), (True, 1)],
        ):
            with self.subTest(pairs=pairs):
                record = hou.trace_record("Supplementary Fig. 18", 1, pairs)
                self.assertEqual(record["status"], "failed")
                self.assertTrue(record["errors"])
                self.assertNotIn("retention", record)
                json.dumps(record, allow_nan=False)

    def test_trailing_empty_rows_counted_but_leading_blank_is_invalid(self):
        record = hou.trace_record("Supplementary Fig. 18", 1, [(0, 2), (1, 2), (None, None), (None, None)])
        self.assertEqual(record["status"], "ok")
        self.assertEqual(record["trailing_fully_empty_rows"], 2)
        self.assertEqual(record["source_rows_seen"], 4)
        self.assertEqual(record["source_coordinates"], "A5:B6")
        leading = hou.trace_record("Supplementary Fig. 18", 1, [(None, None), (0, 2), (1, 2)])
        self.assertEqual(leading["status"], "failed")

    def test_window_coverage_is_required(self):
        with self.assertRaises(ValueError):
            hou.window_mean([0.15, 1], [2, 2], 0.1, 0.2)

    def test_synthetic_workbook_all_records_and_single_sheet_passes(self):
        digest = self.make_workbook()
        calls = []
        original_iter_rows = openpyxl.worksheet._read_only.ReadOnlyWorksheet.iter_rows
        def counted_iter_rows(sheet, *args, **kwargs):
            calls.append(sheet.title)
            return original_iter_rows(sheet, *args, **kwargs)
        with mock.patch.object(openpyxl.worksheet._read_only.ReadOnlyWorksheet, "iter_rows", counted_iter_rows):
            result = hou.analyze_workbook(self.path, digest)
        self.assertEqual(calls, ["Fig. 5e", "Supplementary Fig. 16", "Supplementary Fig. 18"])
        self.assertEqual(result["status"], "ok")
        self.assertTrue(result["all_records_retained"])
        self.assertEqual(result["counts"], {"efficiency_attempted": 3, "traces_attempted": 11,
                                            "efficiency_valid": 3, "traces_valid": 11})
        self.assertEqual(result["source_sheet_passes"], {"Fig. 5e": 1, "Supplementary Fig. 16": 1,
                                                        "Supplementary Fig. 18": 1})
        for index, record in enumerate(result["efficiency"], start=1):
            self.assertEqual(record["healed_original_ratio"], 1.5)
            self.assertEqual(record["difference_native"], index)
        for record in result["traces"]:
            self.assertAlmostEqual(record["retention"], 1 / 17)
            self.assertEqual(record["trailing_fully_empty_rows"], 2)
        self.assertEqual(result["excluded_traces"], [])
        json.dumps(result, allow_nan=False)

    def test_hash_mismatch_never_opens_workbook_and_retains_failures(self):
        self.make_workbook()
        with mock.patch.object(hou.openpyxl, "load_workbook", side_effect=AssertionError("must not open")):
            result = hou.analyze_workbook(self.path, "wrong hash")
        self.assertEqual(result["status"], "failed")
        self.assertTrue(result["all_records_retained"])
        self.assertEqual(result["counts"]["traces_valid"], 0)
        self.assertEqual(result["counts"]["efficiency_valid"], 0)
        self.assertEqual(result["source_sheet_passes"], {})

    def test_wrong_efficiency_label_fails_only_corresponding_pair(self):
        digest = self.make_workbook(lambda wb: setattr(wb["Fig. 5e"]["A6"], "value", "original-2"))
        result = hou.analyze_workbook(self.path, digest)
        self.assertEqual([r["status"] for r in result["efficiency"]], ["ok", "failed", "ok"])
        self.assertNotIn("healed_original_ratio", result["efficiency"][1])
        self.assertTrue(result["all_records_retained"])

    def test_bad_trace_value_and_header_remain_visible(self):
        def change(wb):
            wb["Supplementary Fig. 18"]["H6"] = "bad"
            wb["Supplementary Fig. 16"]["B4"] = "potential"
        digest = self.make_workbook(change)
        result = hou.analyze_workbook(self.path, digest)
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["counts"]["traces_valid"], 9)
        self.assertTrue(result["all_records_retained"])
        failed = [r for r in result["traces"] if r["status"] == "failed"]
        self.assertEqual({(r["sheet"], r["time_column"]) for r in failed},
                         {("Supplementary Fig. 16", "A"), ("Supplementary Fig. 18", "G")})
        self.assertTrue(all("retention" not in record for record in failed))

    def test_missing_sheet_retains_all_eight_attempts(self):
        digest = self.make_workbook(lambda wb: wb.remove(wb["Supplementary Fig. 18"]))
        result = hou.analyze_workbook(self.path, digest)
        self.assertEqual(result["counts"]["traces_valid"], 3)
        self.assertTrue(result["all_records_retained"])

    def test_nonpositive_fe_and_separator_refusal(self):
        digest = self.make_workbook(lambda wb: setattr(wb["Fig. 5e"]["B3"], "value", 0))
        result = hou.analyze_workbook(self.path, digest)
        self.assertEqual(result["counts"]["efficiency_valid"], 2)
        digest = self.make_workbook(lambda wb: setattr(wb["Fig. 5e"]["B5"], "value", 0))
        result = hou.analyze_workbook(self.path, digest)
        self.assertEqual(result["counts"]["efficiency_valid"], 0)
        self.assertTrue(result["all_records_retained"])

    def prepare_cli_inputs(self):
        base = Path(self.directory.name)
        (base / "sources").mkdir()
        for name in ("PROTOCOL.md", "TEST_PROTOCOL.md", "TEST_CLARIFICATION.md"):
            (base / name).write_bytes((hou.BASE / name).read_bytes())
        for name in ("analyze_hou.py", "test_analyze_hou.py", "sources/native-inventory.json",
                     "sources/time-axis-schema.json"):
            (base / name).write_text("synthetic provenance input")
        return base

    def test_cli_provenance_and_frozen_protocols(self):
        base = self.prepare_cli_inputs()
        digest = self.make_workbook()
        self.path.rename(base / "sources" / "hou-source.xlsx")
        original_analyze = hou.analyze_workbook
        with mock.patch.object(hou, "BASE", base), mock.patch.object(
            hou, "analyze_workbook", side_effect=lambda path: original_analyze(path, digest)
        ), mock.patch("builtins.print"):
            self.assertEqual(hou.main(), 0)
        result = json.loads((base / "results" / "hou-test.json").read_text())
        self.assertEqual(result["provenance"]["TEST_PROTOCOL.md"]["sha256"], hou.TEST_PROTOCOL_SHA256)
        self.assertEqual(result["provenance"]["PROTOCOL.md"]["sha256"], hou.PROTOCOL_SHA256)
        self.assertEqual(result["provenance"]["TEST_CLARIFICATION.md"]["sha256"],
                         hou.sha256(base / "TEST_CLARIFICATION.md"))
        self.assertEqual(result["workbook"]["sha256"], digest)
        for item in result["provenance"].values():
            self.assertEqual(item["sha256"], hou.sha256(item["path"]))
        self.assertIn("openpyxl", result["runtime"])

    def test_protocol_change_blocks_before_workbook_access(self):
        base = self.prepare_cli_inputs()
        (base / "TEST_PROTOCOL.md").write_text("changed protocol")
        with mock.patch.object(hou, "BASE", base), mock.patch.object(hou, "analyze_workbook") as analyze:
            with self.assertRaisesRegex(ValueError, "frozen protocol hash mismatch"):
                hou.main()
        analyze.assert_not_called()
        self.assertFalse((base / "results" / "hou-test.json").exists())

    def test_output_is_finite_json_and_cannot_overwrite(self):
        path = Path(self.directory.name) / "results" / "hou-test.json"
        hou.write_result({"synthetic": True}, path)
        with self.assertRaises(FileExistsError):
            hou.write_result({"synthetic": False}, path)
        self.assertEqual(json.loads(path.read_text()), {"synthetic": True})
        with self.assertRaises(ValueError):
            hou.write_result({"invalid": float("nan")}, path.parent / "invalid.json")
        self.assertFalse((path.parent / "invalid.json").exists())


if __name__ == "__main__":
    unittest.main()
