import hashlib
import json
from pathlib import Path, PurePosixPath
import stat
import zipfile

import fitz
import openpyxl

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "sources"


def main():
    output = SOURCE / "native-inventory.json"
    if output.exists():
        raise FileExistsError(output)
    receipt = json.loads((SOURCE / "receipt.json").read_text())
    checks = {}
    for record in receipt["requests"]:
        path = SOURCE / record["name"]
        data = path.read_bytes()
        checks[path.name + "_size"] = len(data) == record["bytes"]
        checks[path.name + "_sha256"] = hashlib.sha256(data).hexdigest() == record["sha256"]
    html = (SOURCE / "hou-main.html").read_text()
    checks["main_doi"] = receipt["resource_doi"] in html
    checks["main_title"] = receipt["source_title"] in html
    checks["main_cc_by_4"] = "creativecommons.org/licenses/by/4.0/" in html
    checks["native_si_signature"] = (SOURCE / "hou-si.pdf").read_bytes().startswith(b"%PDF")
    with fitz.open(SOURCE / "hou-si.pdf") as pdf:
        pages = len(pdf)
        first = pdf[0].get_text()
    checks["si_title"] = "self-healing electrocatalytic system" in first
    checks["si_21_pages"] = pages == 21
    with zipfile.ZipFile(SOURCE / "hou-source.xlsx") as archive:
        members = archive.infolist()
        checks["xlsx_entry_count"] = len(members) < 10000
        checks["xlsx_total_expansion"] = sum(x.file_size for x in members) < 300_000_000
        checks["xlsx_safe_paths"] = all(
            not PurePosixPath(x.filename).is_absolute()
            and ".." not in PurePosixPath(x.filename).parts
            and not stat.S_ISLNK(x.external_attr >> 16) for x in members)
        checks["xlsx_crc"] = archive.testzip() is None
        checks["xlsx_no_vba"] = not any("vbaproject" in x.filename.lower() for x in members)
        expanded = sum(x.file_size for x in members)
    sheets = []
    workbook = openpyxl.load_workbook(SOURCE / "hou-source.xlsx", read_only=True, data_only=False)
    try:
        for sheet in workbook:
            labels = []
            for row in sheet.iter_rows(max_row=min(sheet.max_row, 14), max_col=min(sheet.max_column, 18)):
                for cell in row:
                    if isinstance(cell.value, str):
                        labels.append(dict(cell=cell.coordinate, text=cell.value))
            sheets.append(dict(name=sheet.title, rows=sheet.max_row,
                               columns=sheet.max_column, top_text_labels=labels))
    finally:
        workbook.close()
    result = dict(checks=checks, passed=sum(checks.values()), total=len(checks),
                  si_pages=pages, xlsx_members=len(members), xlsx_expanded_bytes=expanded,
                  sheets=sheets, response_values_selected=False, scientific_validation=False)
    output.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))
    assert all(checks.values()), "Native integrity or identity check failed"


if __name__ == "__main__":
    main()
