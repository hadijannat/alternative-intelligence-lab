import hashlib
import json
import math
from pathlib import Path
import sys

import openpyxl
from openpyxl.utils import get_column_letter


BASE = Path(__file__).resolve().parent
WORKBOOK_SHA256 = "3319d65a2cfc4fa43f1ea78b906b5be2a5987d5efa928f1312beaeb38f843315"
PROTOCOL_SHA256 = "1f3eb5253caf78a8a12e8d0a655c05f47b5907b0f9d083b9687a3920fb5ea70b"
TEST_PROTOCOL_SHA256 = "0aa3857b762953d98358f405e139538efa622f61b0f7bbed7dd1a4f0174dc825"
TRACE_SHEETS = {"Supplementary Fig. 16": 3, "Supplementary Fig. 18": 8}


def sha256(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def finite_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def window_mean(u, magnitude, start, end):
    if not 0 <= start < end <= 1 or u[0] > start or u[-1] < end:
        raise ValueError("incomplete comparison window")
    pieces = []
    for x0, x1, y0, y1 in zip(u, u[1:], magnitude, magnitude[1:]):
        left, right = max(x0, start), min(x1, end)
        if left < right:
            yl = y0 + (y1 - y0) * ((left - x0) / (x1 - x0))
            yr = y0 + (y1 - y0) * ((right - x0) / (x1 - x0))
            pieces.append((right - left) * (yl / 2 + yr / 2))
    result = math.fsum(pieces) / (end - start)
    if not math.isfinite(result):
        raise ValueError("nonfinite integrated mean")
    return result


def trace_record(sheet, column, pairs, source_headers=None):
    time_col = get_column_letter(column)
    current_col = get_column_letter(column + 1)
    record = {
        "sheet": sheet, "time_column": time_col, "current_column": current_col,
        "first_data_row": 5, "source_rows_seen": len(pairs),
        "status": "failed", "errors": [],
        "time_native_unit": "minutes" if sheet == "Supplementary Fig. 16" and column == 1 else "seconds",
        "current_scale": "native, not calibrated across columns",
    }
    if source_headers is not None:
        record["source_headers"] = [str(label) for label in source_headers]
        expected_time = "time " if sheet == "Supplementary Fig. 16" and column == 7 else "time"
        if source_headers != [expected_time, "current density"]:
            record["errors"].append({"reason": "source-label mismatch"})
    end = len(pairs)
    while end and pairs[end - 1] == (None, None):
        end -= 1
    record["trailing_fully_empty_rows"] = len(pairs) - end
    record["data_rows_attempted"] = end
    record["source_coordinates"] = f"{time_col}5:{current_col}{end + 4}" if end else None
    t, current = [], []
    for row, (x, y) in enumerate(pairs[:end], start=5):
        if not finite_number(x) or not finite_number(y):
            record["errors"].append({"row": row, "reason": "nonfinite, missing or nonnumeric pair"})
            continue
        t.append(float(x))
        current.append(float(y))
        if len(t) > 1 and t[-1] <= t[-2]:
            record["errors"].append({"row": row, "reason": "time is not strictly increasing"})
    record["finite_pair_count"] = len(t)
    if len(t) < 2:
        record["errors"].append({"reason": "fewer than two paired observations"})
    if record["errors"]:
        return record
    span = t[-1] - t[0]
    if not math.isfinite(span) or span <= 0:
        record["errors"].append({"reason": "invalid time span"})
        return record
    u = [(x - t[0]) / span for x in t]
    magnitude = [abs(y) for y in current]
    try:
        early = window_mean(u, magnitude, 0.10, 0.20)
        late = window_mean(u, magnitude, 0.90, 1.00)
        if early <= 0:
            raise ValueError("early-window magnitude is not positive")
        retention = late / early
        if not math.isfinite(retention):
            raise ValueError("nonfinite retention")
    except (ValueError, OverflowError, ZeroDivisionError) as exc:
        record["errors"].append({"reason": str(exc)})
        return record
    record.update({
        "status": "ok", "first_time_native": t[0], "last_time_native": t[-1],
        "early_window_native": [t[0] + span * v for v in (0.1, 0.2)],
        "late_window_native": [t[0] + span * v for v in (0.9, 1.0)],
        "early_mean_magnitude_native": early, "late_mean_magnitude_native": late,
        "retention": retention, "fractional_decline": 1 - retention,
        "persistent_early_reference": 1, "reference_residual": retention - 1,
        "nominal_deterioration": retention < 1,
    })
    return record


def efficiency_records(rows):
    records = []
    for index, first in enumerate((3, 6, 9), start=1):
        record = {"sheet": "Fig. 5e", "comparison": index, "status": "failed",
                  "source_coordinates": [f"A{first}:B{first}", f"A{first+1}:B{first+1}"], "errors": []}
        original_label, original = rows[first - 1][:2]
        healed_label, healed = rows[first][:2]
        record["labels"] = [str(original_label), str(healed_label)]
        if (original_label, healed_label) != (f"origina-{index}", f"selfhealing-{index}"):
            record["errors"].append("source-label mismatch")
        if not all(finite_number(v) and v > 0 for v in (original, healed)):
            record["errors"].append("FE must be positive finite numeric")
        if not record["errors"]:
            ratio, difference = healed / original, healed - original
            if not math.isfinite(ratio) or not math.isfinite(difference):
                record["errors"].append("nonfinite FE arithmetic")
            else:
                record.update({"status": "ok", "original_native": original, "healed_native": healed,
                               "healed_original_ratio": ratio, "difference_native": difference,
                               "nominal_recovery": ratio > 1, "units": "unlabeled native FE scale"})
        records.append(record)
    return records


def analyze_workbook(path, expected_sha256=WORKBOOK_SHA256):
    actual_hash = sha256(path)
    result = {"workbook": {"path": str(Path(path).resolve()), "sha256": actual_hash,
                           "expected_sha256": expected_sha256}, "errors": [], "efficiency": [], "traces": [],
              "windows_normalized": {"early": [0.1, 0.2], "late": [0.9, 1.0]},
              "interpolation": "piecewise linear in sampled current magnitude",
              "excluded_efficiency_separator_rows": [5, 8], "excluded_traces": [],
              "source_sheet_passes": {}, "full_integrated_requirements_closed": []}
    failure = None
    workbook = None
    if actual_hash != expected_sha256:
        failure = "workbook-hash mismatch; no workbook parsed"
    else:
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            failure = f"workbook parsing failed: {type(exc).__name__}: {exc}"
    if failure:
        result["errors"].append(failure)
    try:
        for name, count in {"Fig. 5e": 0, **TRACE_SHEETS}.items():
            try:
                if failure:
                    raise ValueError(failure)
                sheet = workbook[name]
                max_col = 2 if count == 0 else 3 * count - 1
                iterator = sheet.iter_rows(min_row=1, max_col=max_col, values_only=True)
                result["source_sheet_passes"][name] = 1
                if count == 0:
                    rows = list(iterator)
                    if len(rows) < 10 or rows[0][0] != "Fig. 5e":
                        raise ValueError("missing Fig. 5e source rows/title")
                    if any(rows[row - 1][1] != "--" for row in (5, 8)):
                        raise ValueError("source separator mismatch")
                    result["efficiency"] = efficiency_records(rows)
                    continue
                headers = [next(iterator) for _ in range(4)]
                if headers[0][0] != f"Fig. S{name.rsplit(' ', 1)[1]}":
                    raise ValueError("source title mismatch")
                columns = [1 + 3 * index for index in range(count)]
                pairs = [[] for _ in columns]
                for row in iterator:
                    for values, column in zip(pairs, columns):
                        values.append((row[column - 1], row[column]))
                for values, column in zip(pairs, columns):
                    labels = [headers[3][column - 1], headers[3][column]]
                    record = trace_record(name, column, values, labels)
                    result["traces"].append(record)
            except Exception as exc:
                reason = f"{type(exc).__name__}: {exc}"
                result["errors"].append({"sheet": name, "reason": reason})
                if count == 0:
                    result["efficiency"] = [{"sheet": name, "comparison": index, "status": "failed", "errors": [reason]}
                                            for index in range(1, 4)]
                else:
                    result["traces"] = [record for record in result["traces"] if record["sheet"] != name]
                    result["traces"].extend({"sheet": name, "time_column": get_column_letter(1 + 3 * index),
                                              "current_column": get_column_letter(2 + 3 * index),
                                              "status": "failed", "errors": [reason]} for index in range(count))
    finally:
        if workbook is not None:
            workbook.close()
    expected = {(name, get_column_letter(1 + 3 * i)) for name, count in TRACE_SHEETS.items() for i in range(count)}
    observed = [(record["sheet"], record["time_column"]) for record in result["traces"]]
    result["all_records_retained"] = len(observed) == 11 and set(observed) == expected and len(result["efficiency"]) == 3
    result["counts"] = {"efficiency_attempted": len(result["efficiency"]), "traces_attempted": len(observed),
                        "efficiency_valid": sum(r["status"] == "ok" for r in result["efficiency"]),
                        "traces_valid": sum(r["status"] == "ok" for r in result["traces"])}
    result["status"] = "ok" if result["all_records_retained"] and not result["errors"] and all(
        r["status"] == "ok" for r in result["efficiency"] + result["traces"]) else "failed"
    return result


def write_result(result, path):
    payload = json.dumps(result, indent=2, allow_nan=False) + "\n"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x") as stream:
        stream.write(payload)


def main():
    output = BASE / "results" / "hou-test.json"
    if output.exists():
        raise FileExistsError(output)
    provenance = {}
    for name in ("PROTOCOL.md", "TEST_PROTOCOL.md", "TEST_CLARIFICATION.md", "analyze_hou.py", "test_analyze_hou.py",
                 "sources/native-inventory.json", "sources/time-axis-schema.json"):
        provenance[name] = {"path": str(BASE / name), "sha256": sha256(BASE / name)}
    for name, expected in (("PROTOCOL.md", PROTOCOL_SHA256), ("TEST_PROTOCOL.md", TEST_PROTOCOL_SHA256)):
        if provenance[name]["sha256"] != expected:
            raise ValueError(f"frozen protocol hash mismatch: {name}")
    result = analyze_workbook(BASE / "sources" / "hou-source.xlsx")
    result["provenance"] = provenance
    result["runtime"] = {"python": sys.version, "openpyxl": openpyxl.__version__}
    write_result(result, output)
    print(json.dumps({"path": str(output), "status": result["status"], "counts": result["counts"]}))
    return 0 if result["status"] == "ok" else 1


if __name__ == "__main__":
    sys.exit(main())
