"""Verify the native package and inventory observations without response fitting."""
import hashlib
import json
from pathlib import Path

import fitz
import openpyxl

ROOT = Path(__file__).resolve().parent


if __name__ == "__main__":
    receipt = json.loads((ROOT / "receipt.json").read_text())
    checks = dict(receipt["checks"])
    for record in receipt["requests"]:
        data = (ROOT / record["name"]).read_bytes()
        checks[record["name"] + "_unchanged"] = (
            len(data) == record["bytes"] and hashlib.sha256(data).hexdigest() == record["sha256"])
    metadata = json.loads((ROOT / "ou-figshare.json").read_text())
    checks["exact_resource_doi_recomputed"] = metadata["resource_doi"] == receipt["resource_doi"]
    checks["dataset_license"] = metadata["license"]["name"] == "CC BY 4.0"
    native = next(f for f in metadata["files"] if f["name"] == "Raw Tabulated Data from Figures.xlsx")
    data = (ROOT / "ou-tabulated.xlsx").read_bytes()
    checks["workbook_metadata_recomputed"] = (
        len(data) == native["size"] and hashlib.md5(data).hexdigest() == native["computed_md5"])
    pages = {}
    for name, expected in (("ou-main.pdf", 12), ("ou-si.pdf", 35)):
        with fitz.open(ROOT / name) as document:
            pages[name] = len(document)
            checks[name + "_pages"] = len(document) == expected
            checks[name + "_title"] = "Cooperative Fe" in document[0].get_text()
            text = "\n".join(f"--- PDF PAGE {i + 1} ---\n" + page.get_text()
                             for i, page in enumerate(document))
            (ROOT / (name + ".txt")).write_text(text)
            for page_number in ([9] if name == "ou-main.pdf" else [24, 25, 28, 29]):
                document[page_number - 1].get_pixmap(matrix=fitz.Matrix(1.4, 1.4)).save(
                    ROOT / (name + f"-page-{page_number}.png"))
    workbook = openpyxl.load_workbook(ROOT / "ou-tabulated.xlsx", read_only=True, data_only=False)
    inventory = []
    for sheet in workbook:
        labels, formulas, numeric_cells = [], 0, 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas += 1
                elif isinstance(cell.value, str):
                    labels.append([cell.coordinate, cell.value])
                elif isinstance(cell.value, (int, float)):
                    numeric_cells += 1
        inventory.append(dict(sheet=sheet.title, rows=sheet.max_row, columns=sheet.max_column,
                              labels=labels, formulas=formulas, numeric_cells=numeric_cells))
    workbook.close()
    checks["twenty_one_sheets"] = len(inventory) == 21
    result = dict(checks=checks, all_pass=all(checks.values()), pdf_pages=pages,
                  workbook_inventory=inventory, scientific_validation=False,
                  response_points_selected_for_analysis=0, empirical_holdout_validation=False)
    (ROOT / "verification.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(dict(all_pass=result["all_pass"], pdf_pages=pages,
                          sheets=len(inventory), failed=[k for k, v in checks.items() if not v]), indent=2))
    assert result["all_pass"], "Source integrity check failed"
