"""Check native source identity and inventory controls without selecting spectra."""
import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import fitz
import openpyxl

ROOT = Path(__file__).resolve().parent
MEMBERS = {
    "si.pdf": "41467_2025_57624_MOESM1_ESM.pdf",
    "movie-descriptions.pdf": "41467_2025_57624_MOESM2_ESM.pdf",
    "peer-review.pdf": "41467_2025_57624_MOESM9_ESM.pdf",
    "source-data.xlsx": "41467_2025_57624_MOESM10_ESM.xlsx",
}


def main():
    receipt = json.loads((ROOT / "receipt.json").read_text())
    checks = {}
    for record in receipt["requests"]:
        data = (ROOT / record["name"]).read_bytes()
        checks[record["name"] + "_unchanged"] = (
            len(data) == record["bytes"]
            and hashlib.sha256(data).hexdigest() == record["sha256"])
    checks["both_requests_successful"] = len(receipt["requests"]) == 2 and all(
        r["status"] == 200 and not r.get("error") for r in receipt["requests"])
    assert all(checks.values()), "Native receipt check failed"
    article = ET.parse(ROOT / "article.xml").getroot()
    ids = {e.get("pub-id-type"): e.text for e in article.findall("./front/article-meta/article-id")}
    checks["exact_doi"] = ids.get("doi") == "10.1038/s41467-025-57624-w"
    checks["exact_pmc"] = ids.get("pmcid") == "PMC11897212"
    license_text = " ".join("".join(e.itertext()) for e in article.findall(".//license"))
    checks["main_license_nc_nd"] = "Attribution-NonCommercial-NoDerivatives 4.0" in license_text
    # Full native text stays local; the exporter excludes source payloads.
    (ROOT / "article.txt").write_text("\n".join(
        f"--- {s.get('id')} ---\n" + " ".join("".join(s.itertext()).split())
        for s in article.findall(".//body/sec")))
    extracted = []
    with zipfile.ZipFile(ROOT / "supplementary.zip") as archive:
        checks["zip_crc"] = archive.testzip() is None
        checks["unique_zip_names"] = len(archive.namelist()) == len(set(archive.namelist()))
        assert checks["zip_crc"] and checks["unique_zip_names"]
        # Read exact trusted member names into fixed local filenames; no extractall.
        for name, member in MEMBERS.items():
            info = archive.getinfo(member)
            assert not info.is_dir() and info.file_size < 20_000_000
            data = archive.read(member)
            target = ROOT / name
            if target.exists():
                assert target.read_bytes() == data, f"Existing native file differs: {name}"
            else:
                target.write_bytes(data)
            extracted.append(dict(name=name, member=member, bytes=len(data),
                                  sha256=hashlib.sha256(data).hexdigest()))
    for name, expected_pages in (("si.pdf", 36), ("peer-review.pdf", 2)):
        with fitz.open(ROOT / name) as document:
            checks[name + "_pages"] = len(document) == expected_pages
            (ROOT / (name + ".txt")).write_text("\n".join(
                f"--- PDF PAGE {i + 1} ---\n" + page.get_text()
                for i, page in enumerate(document)))
            if name == "si.pdf":
                for number in (7, 8, 9, 10, 12, 13):
                    document[number - 1].get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(
                        ROOT / f"si-page-{number}.png")
    workbook = openpyxl.load_workbook(ROOT / "source-data.xlsx", read_only=True, data_only=False)
    sheets = []
    try:
        for sheet in workbook:
            headers = [{"row": row, "column": column, "text": value}
                       for row, values in enumerate(sheet.iter_rows(max_row=3, values_only=True), 1)
                       for column, value in enumerate(values, 1) if isinstance(value, str)]
            sheets.append(dict(name=sheet.title, rows=sheet.max_row, columns=sheet.max_column,
                               top_three_row_text_labels=headers))
    finally:
        workbook.close()
    checks["workbook_has_figure3h"] = any(s["name"] == "Figure 3h" for s in sheets)
    result = dict(checks=checks, all_pass=all(checks.values()), extracted=extracted,
                  workbook_structure=sheets, protocol_sha256=hashlib.sha256(
                      (ROOT.parent / "PROTOCOL.md").read_bytes()).hexdigest(),
                  response_points_selected_for_analysis=0, numerical_response_fit=False,
                  empirical_holdout_validation=False, scientific_validation=False)
    (ROOT / "verification.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(dict(all_pass=result["all_pass"], checks=len(checks), sheets=[s["name"] for s in sheets],
                          failed=[k for k, v in checks.items() if not v]), indent=2))
    assert result["all_pass"], "Inspect retained verification result"


if __name__ == "__main__":
    main()
